"""FE — 평가·리포트 (STEP 8).

정확도 평가는 실측(source_flag == 'measured') 데이터만 사용한다 (규칙 R3).
모든 리포트에 모델 버전과 학습 Lot 수를 기록한다 (NFR-08, FE-06) —
성능 변화가 모델 개선 때문인지 데이터 증가 때문인지 구분하기 위함이다.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field

import numpy as np
import pandas as pd

from dry_process_ai.config import EVALUATION_DIR, STAGES
from dry_process_ai.core.infer.predictor import Predictor
from dry_process_ai.core.model.builder import (
    AUX_OUTPUT_COLUMNS,
    FINAL_OUTPUT_COLUMNS,
    PERF_OUTPUT_COLUMNS,
    STAGE_OUTPUT_COLUMNS,
)
from dry_process_ai.rules import constraints

# B등급 평가 유보 기준 건수 (FE-03)
B_GRADE_MIN_COUNT = 5


@dataclass
class EvaluationReport:
    model_version: str
    train_lot_count: int
    measured_lot_count: int
    a_grade: dict = field(default_factory=dict)          # FE-01
    per_stage: dict = field(default_factory=dict)        # FE-02
    b_grade: dict = field(default_factory=dict)          # FE-03
    consistency: dict = field(default_factory=dict)      # FE-04
    field_match: dict = field(default_factory=dict)      # FE-05
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "model_version": self.model_version,
            "train_lot_count": self.train_lot_count,
            "measured_lot_count": self.measured_lot_count,
            "a_grade": self.a_grade,
            "per_stage": self.per_stage,
            "b_grade": self.b_grade,
            "consistency": self.consistency,
            "field_match": self.field_match,
            "notes": self.notes,
        }


def _metrics(pred: pd.Series, actual: pd.Series) -> dict | None:
    mask = actual.notna() & pred.notna()
    if not mask.any():
        return None
    p, a = pred[mask].to_numpy(dtype=float), actual[mask].to_numpy(dtype=float)
    rmse = float(np.sqrt(np.mean((p - a) ** 2)))
    nonzero = np.abs(a) > 1e-12
    mape = float(np.mean(np.abs((p[nonzero] - a[nonzero]) / a[nonzero])) * 100.0) if nonzero.any() else None
    return {"rmse": rmse, "mape_pct": mape, "n": int(mask.sum())}


def evaluate_model(predictor: Predictor, df: pd.DataFrame, train_lot_count: int) -> EvaluationReport:
    """FE-01~04 통합 평가. df 는 fetch_dataset() 형식 (source_flag 필수)."""
    if "source_flag" not in df.columns:
        raise ValueError("source_flag 컬럼이 필요하다 — 생성 데이터 혼입 차단 (규칙 R3)")
    measured = df[df["source_flag"] == "measured"].copy()

    report = EvaluationReport(
        model_version=predictor.model_version,
        train_lot_count=train_lot_count,
        measured_lot_count=len(measured),
    )
    if measured.empty:
        report.notes.append("실측 Lot 이 없어 정확도 평가를 유보한다 — 워크플로우·물리 제약 준수가 판정 기준")
        return report

    pred = predictor.predict_frame(measured)

    # FE-01: A등급 정확도 (NFR-03: 두께 MAPE < 10 %, 합제밀도 MAPE < 5 %)
    for col in FINAL_OUTPUT_COLUMNS:
        m = _metrics(pred[col], measured[col])
        if m:
            report.a_grade[col] = m
    if len(measured) < 30:
        report.notes.append(
            f"실측 {len(measured)}건 < 30건 — NFR-03 정확도 목표 적용 유예 구간 (제약 상황 13장)"
        )

    # FE-02: 단계별 정확도 — 취약 단계 식별
    for stage in STAGES:
        stage_metrics = {}
        for target in ("composite_thickness_um", "composite_density_gcc", "areal_capacity_mah_cm2"):
            col = f"{stage}_{target}"
            m = _metrics(pred[col], measured[col]) if col in measured.columns else None
            if m:
                stage_metrics[target] = m
        if stage_metrics:
            report.per_stage[stage] = stage_metrics

    # FE-03: B등급 조건부 정확도 — 실측 건수 병기, 미달 시 유보
    for col in AUX_OUTPUT_COLUMNS:
        n = int(measured[col].notna().sum()) if col in measured.columns else 0
        if n >= B_GRADE_MIN_COUNT:
            report.b_grade[col] = {**_metrics(pred[col], measured[col]), "measured_count": n}
        else:
            report.b_grade[col] = {"measured_count": n, "status": f"실측 {n}건 < {B_GRADE_MIN_COUNT}건 — 평가 유보"}

    # FE-04: 정합성·제약 준수율 (목표: 정합률 100 %, 위반율 0 %)
    # 사용자에게 제시되는 최종 출력과 동일하게, 용량·로딩은 두께·밀도·조성으로
    # 정합화된 값을 검사한다 (predictor.predict_one 과 동일 경로 — 규칙 R5).
    from dry_process_ai.rules import physics

    total, violation_free = 0, 0
    for idx in pred.index:
        am_raw = measured.loc[idx, "active_material_content"]
        am_frac = float(am_raw) / 100.0 if pd.notna(am_raw) else None
        stage_values = {}
        for s in STAGES:
            t = float(pred.loc[idx, f"{s}_composite_thickness_um"])
            d = float(pred.loc[idx, f"{s}_composite_density_gcc"])
            loading = physics.loading_mg_cm2(t, d)
            stage_values[s] = {
                "composite_thickness_um": t,
                "composite_density_gcc": d,
                "loading_mg_cm2": loading,
                "areal_capacity_mah_cm2": (
                    physics.areal_capacity_mah_cm2(loading, am_frac) if am_frac is not None
                    else float(pred.loc[idx, f"{s}_areal_capacity_mah_cm2"])
                ),
            }
        rep = constraints.check_stage_sequence(stage_values, active_material_fraction=am_frac)
        total += 1
        if rep.passed:
            violation_free += 1
    report.consistency = {
        "checked": total,
        "violation_free": violation_free,
        "compliance_rate_pct": round(violation_free / total * 100.0, 2) if total else None,
        "violation_rate_pct": round((1 - violation_free / total) * 100.0, 2) if total else None,
    }
    return report


def field_match_rate(
    predicted_intervals: list[tuple[float, float]],
    measured_values: list[float],
) -> dict:
    """FE-05 — 실증 일치도: 실측이 예측 구간 내에 든 비율 (목표 3/5 이상)."""
    hits = sum(1 for (lo, hi), v in zip(predicted_intervals, measured_values) if lo <= v <= hi)
    n = len(measured_values)
    return {"n": n, "hits": hits, "rate": hits / n if n else None, "target": "3/5 이상"}


def save_report(report: EvaluationReport, directory=EVALUATION_DIR) -> str:
    """eval_{model_version}.json 규약으로 저장 (FE-06, 12.2)."""
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"eval_{report.model_version}.json"
    path.write_text(json.dumps(report.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def save_report_docx(report: EvaluationReport, directory=EVALUATION_DIR) -> str:
    """평가 리포트 문서 산출 (python-docx, FE-06)."""
    from docx import Document

    directory.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("건식 공정 최적화 예측 프로그램 — 종합 평가 리포트", level=1)
    doc.add_paragraph(f"모델 버전: {report.model_version}")
    doc.add_paragraph(f"학습 Lot 수: {report.train_lot_count} (실측 {report.measured_lot_count}건)")

    doc.add_heading("A등급 정확도 (FE-01)", level=2)
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "변수", "RMSE", "MAPE (%)", "N"
    for name, m in report.a_grade.items():
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = f"{m['rmse']:.4g}"
        cells[2].text = f"{m['mape_pct']:.2f}" if m.get("mape_pct") is not None else "-"
        cells[3].text = str(m["n"])

    doc.add_heading("정합성·제약 준수율 (FE-04)", level=2)
    doc.add_paragraph(json.dumps(report.consistency, ensure_ascii=False))

    for note in report.notes:
        doc.add_paragraph(f"※ {note}")

    path = directory / f"eval_{report.model_version}.docx"
    doc.save(str(path))
    return str(path)


def closed_loop_feedback(report: EvaluationReport) -> dict:
    """FE-07 — 폐루프 반영: 정확도 낮은 목표 변수는 FR 추천 우선순위 상향 대상,
    제약 위반 구간은 FP 이상치 규칙·FT 출력 제약 반영 대상으로 식별한다."""
    weak_targets = []
    for name, m in report.a_grade.items():
        if m.get("mape_pct") is not None:
            threshold = 5.0 if "density" in name else 10.0
            if m["mape_pct"] > threshold:
                weak_targets.append(name)
    for stage, metrics in report.per_stage.items():
        for target, m in metrics.items():
            if m.get("mape_pct") is not None and m["mape_pct"] > 10.0:
                weak_targets.append(f"{stage}_{target}")
    return {
        "raise_recommendation_priority": weak_targets,
        "constraint_review_needed": (report.consistency.get("violation_rate_pct") or 0) > 0,
    }
