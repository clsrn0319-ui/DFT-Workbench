"""배치 스크리닝 — 수백 개 바인더 후보의 일괄 DFT 계산과 적합/부적합 자동 판정.

기획서(05_DFT-Workbench_배치스크리닝_기획서) 구현체. 세 부분으로 이루어진다.

  1. 캠페인 저장소 — 후보 목록 + 계산 조건 + 판정 기준 + 진행 상태를 하나로
     묶은 실행 단위. data/campaigns.json 에 지속 저장되어 서버 재시작 후에도
     이어서 진행된다 (체크포인트-재개).
  2. 깔때기 오케스트레이터 — 단계(빠름 → 표준 → 정밀)별로 작업을 제출하고,
     완료를 감지해 통과자를 골라 다음 단계를 자동 제출하는 백그라운드 스레드.
     기존 작업 큐·엔진을 그대로 쓰며, 배치 작업은 단건보다 낮은 우선순위로
     들어가 일상 사용을 방해하지 않는다.
  3. 판정 엔진 — 산화·환원 전위를 활물질 구동 «범위»와 대조해
     적합 / 조건부 / 부적합 / 판정 불가 등급을 부여한다. 판정식은 기존
     esw.containment 와 같은 포함 관계이며, 안정성 마진만 더해졌다.

판정은 열역학적 스크리닝이다 — «분해될 수 있는가»를 보는 것이지
«실제 배터리에서 못 쓴다»는 뜻이 아니다 (EC/SEI 반례). 이 문구는 화면에도
그대로 표시된다.
"""

import json
import os
import threading
import time
import uuid
from pathlib import Path

from rdkit import Chem

from . import confsens, esw, geometry, presets, scoring, store, worker

DATA_DIR = Path(os.environ.get(
    "RHOBENCH_DATA_DIR", Path(__file__).resolve().parent.parent / "data"))
CAMPAIGNS_FILE = DATA_DIR / "campaigns.json"

# 자원 보호 — 환경변수로 조정
MAX_CANDIDATES = int(os.environ.get("RHOBENCH_MAX_BATCH", "500"))
BATCH_PARALLEL = max(1, int(os.environ.get("RHOBENCH_BATCH_PARALLEL", "1")))
# 실패 누적률이 이 값을 넘으면 캠페인을 자동 일시정지하고 관리자 확인을 유도
FAIL_PAUSE_RATIO = float(os.environ.get("RHOBENCH_BATCH_FAIL_RATIO", "0.3"))
# 수직(빠름) 전위는 구조 완화가 빠져 EA를 과소평가 — 환원 위험을 낮잡는다.
# 깔때기 «컷 순위»에서만 환원 전위를 이만큼 올려 보수적으로 비교한다.
VERTICAL_RED_BUFFER = float(os.environ.get("RHOBENCH_VERTICAL_RED_BUFFER", "0.5"))
# 배치 작업 결과 슬림화 — 전자밀도 구름 등 무거운 시각화 데이터를 지워
# 수백 개 캠페인에서 jobs.json 크기·로딩 시간을 줄인다 ("0"으로 끄기)
SLIM_BATCH = os.environ.get("RHOBENCH_BATCH_SLIM", "1") != "0"
# 단계 간 구조 승계 — 이전 단계에서 확정한 구조를 다음 단계의 초기 구조로 넘겨
# conformer 탐색·DFT 재순위화를 건너뛴다 (Boltzmann 앙상블 단계는 제외, "0"으로 끄기)
CHAIN_GEOMETRY = os.environ.get("RHOBENCH_CHAIN_GEOMETRY", "1") != "0"

RESTART_ERROR = "서버 재시작으로 중단됨"

# 분자당 대략 소요 시간(초) — 실행 확인 화면의 추정용. 실측이 쌓이면 실측 우선.
ESTIMATE_S = {"빠름": 120, "표준": 1200, "정밀": 3600}

# 판정에 필요한 목적 — 전위가 없으면 판정 자체가 불가능하다
SCREEN_PURPOSE = "전자구조 + 산화/환원 전위"

THERMO_NOTE = ("이 판정은 열역학적 스크리닝입니다 — «분해될 수 있는가»를 보는 "
               "것이지 «실제 배터리에서 못 쓴다»는 뜻이 아닙니다. 대표 반례: "
               "EC는 계산상 음극에서 환원 분해되지만 그 분해 산물이 안정적인 "
               "보호막(SEI)을 만들어 상용 전해액의 표준 성분으로 쓰입니다.")

_lock = threading.RLock()
_campaigns: dict[str, dict] = {}
_loaded = False


# ---------------------------------------------------------------- 저장소
def _load():
    global _loaded
    if _loaded:
        return
    _loaded = True
    if CAMPAIGNS_FILE.exists():
        try:
            for c in json.loads(CAMPAIGNS_FILE.read_text(encoding="utf-8")):
                _campaigns[c["id"]] = c
        except (json.JSONDecodeError, OSError):
            pass


def _persist():
    DATA_DIR.mkdir(exist_ok=True)
    tmp = CAMPAIGNS_FILE.with_suffix(".tmp")
    tmp.write_text(json.dumps(list(_campaigns.values()), ensure_ascii=False, indent=1),
                   encoding="utf-8")
    tmp.replace(CAMPAIGNS_FILE)


def _log(camp: dict, msg: str):
    camp.setdefault("logs", []).append(f"[{time.strftime('%m-%d %H:%M:%S')}] {msg}")


# 작용기 태그 (기획서 9.1) — 표·필터용 표시. 위에서 아래로 검사하며 겹침 허용.
FG_SMARTS = [
    ("–COOH", "C(=O)[OX2H1]"),
    ("–SO₃H", "S(=O)(=O)[OX2H]"),
    ("에스터", "C(=O)O[#6]"),
    ("아마이드", "C(=O)[NX3]"),
    ("–CN", "[NX1]#[CX2]"),
    ("–NH₂", "[NX3;H2;!$(NC=O)]"),
    ("–OH", "[OX2H;!$([OX2H]C=O)]"),
    ("–F", "[F]"),
    ("방향족", "a"),
    ("C=C", "[CX3]=[CX3]"),
]
_FG_PATTERNS = [(name, Chem.MolFromSmarts(sm)) for name, sm in FG_SMARTS]


def _fgroups(smiles: str) -> list[str]:
    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        return []
    return [name for name, patt in _FG_PATTERNS
            if patt is not None and mol.HasSubstructMatch(patt)]


# ---------------------------------------------------------------- 후보 파싱
def _canonical(smiles: str):
    mol = Chem.MolFromSmiles(smiles)
    return Chem.MolToSmiles(mol) if mol is not None else None


def _split_line(line: str) -> list[str]:
    for sep in (",", "\t", ";"):
        if sep in line:
            return [f.strip().strip('"') for f in line.split(sep)]
    return [line.strip()]


def parse_candidates(text: str, structure: str = "모노머",
                     max_atoms: int = 60) -> dict:
    """CSV(이름,SMILES[,메모]) 또는 한 줄 SMILES 목록을 검증된 후보로 바꾼다.

    - 열 순서는 자동 감지 — SMILES 로 해석되는 필드를 구조로, 나머지를 이름으로
    - 헤더 행(smiles/name 등 단어만 있고 구조가 없는 행)은 자동 건너뜀
    - 중복은 캐노니컬 SMILES 기준으로 제거 (첫 등장만 유지)
    - 2량체·3량체 구조 선택 시 올리고머로 전개한 뒤 원자 수 상한을 검사
    """
    n_units = {"모노머": 1, "2량체": 2, "3량체": 3}.get(structure, 1)
    rows, seen = [], {}
    ok_count = 0
    for line_no, raw in enumerate(text.splitlines(), start=1):
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        fields = _split_line(line)
        smiles, name = None, None
        for f in fields:
            if smiles is None and f and _canonical(f) is not None:
                smiles = f
            elif f and name is None:
                name = f
        row = {"line": line_no, "name": name or smiles or line,
               "smiles": smiles, "ok": False, "error": None}
        if smiles is None:
            # 헤더 행이면 조용히 건너뛴다 — 오류 목록에 잡음을 넣지 않는다
            joined = " ".join(fields).lower()
            if line_no <= 2 and any(w in joined for w in ("smiles", "name", "이름", "구조")):
                continue
            row["error"] = "SMILES 로 해석되는 필드가 없습니다"
            rows.append(row)
            continue
        canonical = _canonical(smiles)
        if canonical in seen:
            row["error"] = f"{seen[canonical]}행과 중복 구조 (제외)"
            rows.append(row)
            continue
        seen[canonical] = line_no
        row["canonical"] = canonical
        row["fgroups"] = _fgroups(canonical)
        calc_smiles = smiles
        if n_units > 1:
            try:
                calc_smiles = geometry.oligomerize(smiles, n_units)
            except geometry.GeometryError as exc:
                row["error"] = f"{structure} 전개 실패: {exc}"
                rows.append(row)
                continue
        atoms = geometry.atom_count(calc_smiles)
        row.update({"calc_smiles": calc_smiles, "atoms": atoms})
        if atoms > max_atoms:
            row["error"] = f"원자 수 {atoms}개 — 상한 {max_atoms}개 초과"
            rows.append(row)
            continue
        row["ok"] = True
        ok_count += 1
        rows.append(row)
    return {"rows": rows, "n_ok": ok_count,
            "n_error": sum(1 for r in rows if not r["ok"]),
            "structure": structure, "max_atoms": max_atoms}


def xlsx_to_text(data: bytes) -> str:
    """엑셀(.xlsx) 첫 시트를 CSV 텍스트로 바꿔 parse_candidates 에 그대로 넘긴다."""
    import io as _io

    from openpyxl import load_workbook
    wb = load_workbook(_io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        lines = []
        for row in ws.iter_rows(max_row=MAX_CANDIDATES + 50, values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(",".join(cells))
    finally:
        wb.close()
    return "\n".join(lines)


# ---------------------------------------------------------------- 판정 엔진
def _potentials(desc: dict):
    """ΔG 기반 전위가 있으면 우선 사용 — esw.diagnose 와 같은 규칙.

    키가 None 값으로 «존재»하는 경우까지 폴백해야 한다 (esw.first_present 참조).
    """
    red = esw.first_present(desc, "reduction_potential_gibbs_v",
                            "reduction_potential_v")
    ox = esw.first_present(desc, "oxidation_potential_gibbs_v",
                           "oxidation_potential_v")
    return red, ox


def judge(desc: dict, electrodes: list[str], margin_v: float,
          windows: dict | None = None, e_abs: float = 1.44) -> dict:
    """전극 구동 범위 전체가 ESW 안에 얼마나 여유 있게 들어오는지로 등급을 매긴다.

    여유(margin) = min(구동 하단 − 환원 전위, 산화 전위 − 구동 상단)
      여유 ≥ 안정성 마진  → 적합
      0 ≤ 여유 < 마진     → 조건부 적합 (계산 오차 범위 — 정밀 재확인 권장)
      여유 < 0            → 부적합 (구동 범위 안에서 열역학적 분해)
    """
    red, ox = _potentials(desc)
    basis = ("ΔG 기반" if desc.get("reduction_potential_gibbs_v") is not None
             else "단열" if desc.get("ea_adiabatic_ev") is not None
             or desc.get("ip_adiabatic_ev") is not None else "수직")
    if red is None or ox is None:
        return {"grade": "판정 불가", "per_electrode": [], "worst_margin_v": None,
                "reduction_v": red, "oxidation_v": ox, "basis": None,
                "note": "산화·환원 전위가 없어 판정할 수 없습니다"}

    # conformer 민감도 (v2.0 P0-5). 편차가 0.20 V 이상이면 단일값이 아니라
    # «범위»로 판정한다 — 범위 전체가 여유를 넘어야 적합, 범위 전체가 침범해야
    # 부적합, 걸치면 조건부. 그 아래 편차는 신뢰도 축(scoring)에서 다룬다.
    sens = desc.get("conformer_sensitivity") or {}
    rng = confsens.judgement_range(sens)
    conformer_range = None
    if rng:
        conformer_range = {
            "reduction_v": [round(red + rng["red"][0], 3), round(red + rng["red"][1], 3)],
            "oxidation_v": [round(ox + rng["ox"][0], 3), round(ox + rng["ox"][1], 3)],
            "spread_v": rng["spread_v"], "n_conformers": rng["n_conformers"],
        }
    per = []
    for key in electrodes:
        win = (windows or esw.ELECTRODE_BY_KEY)[key]
        if conformer_range:
            r_lo, r_hi = conformer_range["reduction_v"]
            o_lo, o_hi = conformer_range["oxidation_v"]
            m_worst = round(min(win["low"] - r_hi, o_lo - win["high"]), 3)
            m_best = round(min(win["low"] - r_lo, o_hi - win["high"]), 3)
            grade = ("적합" if m_worst >= margin_v
                     else "부적합" if m_best < 0 else "조건부")
            per.append({"electrode": key, "label": win["label"], "side": win["side"],
                        "window": [win["low"], win["high"]],
                        "margin_v": m_worst, "margin_best_v": m_best,
                        "grade": grade, "range_judged": True})
            continue
        m = round(min(win["low"] - red, ox - win["high"]), 3)
        grade = "적합" if m >= margin_v else ("조건부" if m >= 0 else "부적합")
        per.append({"electrode": key, "label": win["label"], "side": win["side"],
                    "window": [win["low"], win["high"]],
                    "margin_v": m, "grade": grade})
    worst = min(p["margin_v"] for p in per)
    if all(p["grade"] == "적합" for p in per):
        grade = "적합"
    elif any(p["grade"] == "부적합" for p in per):
        grade = "부적합"
    else:
        grade = "조건부"
    note = None
    if basis == "수직":
        note = ("수직 전위 기반 — 구조 완화·열보정이 빠져 환원 위험을 낮잡을 수 "
                "있습니다. 표준·정밀 단계 재확인을 권장합니다.")
    if conformer_range:
        rn = (f"conformer 편차 σ {conformer_range['spread_v']:.2f} V ≥ 0.20 V — "
              f"{conformer_range['n_conformers']}개 conformer 의 전위 범위 "
              f"(환원 {conformer_range['reduction_v'][0]:+.2f}~{conformer_range['reduction_v'][1]:+.2f} V)로 "
              "판정했습니다. 범위 전체가 여유를 넘어야 적합입니다.")
        note = rn if note is None else note + " " + rn
    conformer_spread = None
    if sens.get("spread_v") is not None:
        conformer_spread = {"spread_v": sens["spread_v"], "rule": sens.get("rule"),
                            "n_conformers": sens.get("n_conformers")}

    # LUMO 불일치 경고 — 수직 EA 판정은 «안전»인데 LUMO 소박 추정
    # (−LUMO − E_abs)은 구동 범위를 침범하는 경우. 수직 EA는 음이온의 구조
    # 완화·평형 용매화가 빠져 환원 위험을 낮잡을 수 있고, 강한 전자 수용체
    # (예: 말레산무수물)에서는 LUMO 쪽이 실제에 가깝다.
    lumo_check = None
    if basis == "수직" and desc.get("lumo_ev") is not None:
        naive_red = round(-desc["lumo_ev"] - e_abs, 3)
        wins_map = windows or esw.ELECTRODE_BY_KEY
        mismatch = [wins_map[k]["label"] for k in electrodes
                    if red <= wins_map[k]["low"] < naive_red]
        if mismatch:
            lumo_check = {
                "naive_red_v": naive_red,
                "mismatch": mismatch,
                "note": (f"수직 EA 환원 전위({red:+.2f} V)는 안전으로 나오지만 "
                         f"LUMO 추정({naive_red:+.2f} V)은 {', '.join(mismatch)} "
                         "구동 범위를 침범합니다 — 수직 계산이 환원 위험을 낮잡았을 "
                         "수 있으니 표준(단열·ΔG) 재계산으로 확정하세요."),
            }
            # 경고에 그치지 않는다 — 상반된 두 지표 중 어느 쪽도 확정이 아니므로,
            # 침범된 활물질의 «적합»을 «조건부(재검증 필요)»로 강등해 적합군에
            # 오르지 못하게 한다. 확정은 표준(단열·ΔG) 재계산이 한다.
            demoted = False
            for p in per:
                if p["label"] in mismatch and p["grade"] == "적합":
                    p["grade"] = "조건부"
                    p["lumo_demoted"] = True
                    demoted = True
            if demoted:
                if any(p["grade"] == "부적합" for p in per):
                    grade = "부적합"
                elif all(p["grade"] == "적합" for p in per):
                    grade = "적합"
                else:
                    grade = "조건부"
                lumo_check["demoted"] = True
                lumo_check["note"] = ("적합 등급을 «조건부(재검증 필요)»로 강등 — "
                                      + lumo_check["note"])
    return {"grade": grade, "per_electrode": per, "worst_margin_v": worst,
            "reduction_v": red, "oxidation_v": ox, "basis": basis, "note": note,
            "lumo_check": lumo_check,
            "conformer_range": conformer_range, "conformer_spread": conformer_spread}


# ---------------------------------------------------------------- 캠페인
def _stage_settings(camp: dict, stage_idx: int) -> dict:
    s = dict(camp["settings"])
    s["accuracy"] = camp["stages"][stage_idx]["accuracy"]
    s["purpose"] = SCREEN_PURPOSE
    # 5대 Score 캠페인은 마지막 단계에서 물성 지문(접착·Li⁺·용매화·BDE)까지
    # 계산한다 — 그 전 단계는 전위만으로 깔때기를 돌려 비용을 아낀다 (기획서 11.1)
    if ((camp.get("scoring") or {}).get("enabled")
            and stage_idx == len(camp["stages"]) - 1):
        s["purpose"] = "물성 지문 (확장 기술자 전체)"
    # 후보는 이미 올리고머로 전개되어 있다 — 엔진이 다시 전개하지 않도록
    s["structure"] = "모노머"
    return s


def create_campaign(name: str, candidates: list[dict], electrodes: list[str],
                    margin_v: float, stages: list[dict], settings: dict,
                    custom_electrodes: list | None = None,
                    scoring_cfg: dict | None = None) -> dict:
    """검증된 후보 목록으로 캠페인을 만들고 즉시 시작한다."""
    with _lock:
        _load()
        cid = "CAM-" + time.strftime("%Y%m%d") + "-" + uuid.uuid4().hex[:6].upper()
        # 동시 실행은 1개 — 이미 도는 캠페인이 있으면 대기열(QUEUED)로 등록되고,
        # 앞 캠페인이 끝나는 대로 오케스트레이터가 자동 시작한다
        busy = any(c["status"] == "RUNNING" for c in _campaigns.values())
        camp = {
            "id": cid,
            "name": name or cid,
            "status": "QUEUED" if busy else "RUNNING",
            # 상태: QUEUED · RUNNING · PAUSED · DONE · CANCELLED
            "createdAt": time.time(),
            "finishedAt": None,
            "electrodes": electrodes,
            "customElectrodes": custom_electrodes or [],
            "scoring": scoring_cfg,          # 5대 Score 설정 (기획서 7장)
            "protocol": scoring.PROTOCOL_VERSION,
            "margin_v": margin_v,
            "stages": [{"accuracy": st["accuracy"], "keep": st.get("keep"),
                        "threshold": st.get("threshold")}
                       for st in stages],
            "stageIndex": 0,
            "settings": settings,
            "structure": settings.get("structure", "모노머"),
            "autoPaused": False,
            "candidates": [
                {"idx": i, "name": c["name"], "smiles": c["smiles"],
                 "canonical": c.get("canonical") or c["smiles"],
                 "calcSmiles": c.get("calc_smiles") or c["smiles"],
                 "atoms": c.get("atoms"),
                 "geometry": c.get("geometry"),   # 업로드 3D 구조 (선택)
                 "fgroups": c.get("fgroups") or [],
                 "alive": True, "failed": False, "cutStage": None,
                 "error": None, "retries": {}, "jobs": {}, "verdict": None}
                for i, c in enumerate(candidates)],
            "logs": [],
        }
        _log(camp, f"캠페인 생성 — 후보 {len(candidates)}개 · "
                   f"단계 {' → '.join(st['accuracy'] for st in camp['stages'])} · "
                   f"전극 {', '.join(electrodes)} · 마진 {margin_v} V")
        if camp["status"] == "QUEUED":
            _log(camp, "다른 캠페인이 진행 중 — 대기열에 등록되어 끝나는 대로 자동 시작합니다")
        _campaigns[cid] = camp
        _persist()
    ensure_started()
    return camp


def get_campaign(cid: str):
    with _lock:
        _load()
        return _campaigns.get(cid)


def list_campaigns():
    with _lock:
        _load()
        return sorted(_campaigns.values(), key=lambda c: c["createdAt"], reverse=True)


def set_status(cid: str, status: str) -> bool:
    """pause / resume / cancel 상태 전이."""
    with _lock:
        _load()
        camp = _campaigns.get(cid)
        if camp is None or camp["status"] in ("DONE", "CANCELLED"):
            return False
        if status == "PAUSED" and camp["status"] != "RUNNING":
            return False   # 대기 중 캠페인은 취소만 가능
        if status == "CANCELLED":
            # 진행 중인 배치 작업도 함께 취소한다
            for c in camp["candidates"]:
                jid = c["jobs"].get(str(camp["stageIndex"]))
                if jid:
                    job = store.get_job(jid)
                    if job and job["status"] in ("QUEUED", "RUNNING"):
                        store.request_cancel(jid)
            camp["finishedAt"] = time.time()
        if status == "RUNNING":
            camp["autoPaused"] = False
        camp["status"] = status
        _log(camp, {"RUNNING": "재개", "PAUSED": "일시정지", "CANCELLED": "취소"}
             .get(status, status))
        _persist()
        return True


def delete_campaign(cid: str) -> bool:
    with _lock:
        _load()
        camp = _campaigns.get(cid)
        if camp is None:
            return False
        if camp["status"] in ("RUNNING", "PAUSED"):
            return False   # 먼저 취소해야 지울 수 있다
        del _campaigns[cid]
        _persist()
        return True


# ---------------------------------------------------------------- 캐시 재사용
def _relevant(settings: dict) -> tuple:
    """결과 재사용 판단에 쓰는 조건 지문 — 이 조합이 같으면 같은 계산이다."""
    exp = settings.get("expert") or {}
    return (settings.get("accuracy"), settings.get("solventId"),
            settings.get("customMixedSolvent") is None or json.dumps(
                settings.get("customMixedSolvent"), sort_keys=True),
            round(float(settings.get("temperature", 298.15)), 2),
            settings.get("referenceElectrode"), settings.get("envType"),
            exp.get("functional"), exp.get("basis"),
            exp.get("charge", 0), exp.get("multiplicity", 1),
            exp.get("conformerSensitivity"))


def _find_cached(canonical: str, settings: dict):
    """같은 구조·같은 조건의 PUBLISHED 결과가 이미 있으면 재계산하지 않는다."""
    want = _relevant(settings)
    for j in store.list_jobs():
        if j["status"] != "PUBLISHED" or not j.get("result"):
            continue
        js = j.get("settings") or {}
        # 전위가 산출된 목적이어야 판정에 쓸 수 있다
        if not any(w in (js.get("purpose") or "") for w in ("전위", "지문", "바인더")):
            continue
        if _relevant(js) != want:
            continue
        if _canonical(j["material"].get("smiles", "")) == canonical:
            return j
    return None


# ---------------------------------------------------------------- 오케스트레이터
def _batch_active_count() -> int:
    return sum(1 for j in store.list_jobs()
               if j.get("campaign") and j["status"] in ("QUEUED", "RUNNING"))


def _e_abs(camp: dict) -> float:
    """캠페인 기준 전극의 절대 전위 — LUMO 소박 추정 환산용."""
    ref = (camp.get("settings") or {}).get("referenceElectrode")
    return presets.ABSOLUTE_POTENTIALS.get(ref, 1.44)


def _windows(camp: dict) -> dict:
    """판정에 쓰는 전극 구동 범위 — 내장 활물질 + 캠페인의 사용자 정의 활물질."""
    wins = dict(esw.ELECTRODE_BY_KEY)
    for ce in camp.get("customElectrodes") or []:
        wins[ce["key"]] = ce
    return wins


def _xyz_atoms(xyz: str):
    """결과의 XYZ 블록 → [[원소, x, y, z], ...] (승계용). 형식이 어긋나면 None."""
    try:
        lines = xyz.strip().splitlines()
        n = int(lines[0].split()[0])
        atoms = []
        for ln in lines[2:2 + n]:
            t = ln.split()
            atoms.append([t[0], float(t[1]), float(t[2]), float(t[3])])
        return atoms if len(atoms) == n else None
    except (ValueError, IndexError):
        return None


def _chained_geometry(c: dict, camp: dict, stage_idx: int):
    """이전 단계의 최적 구조를 이번 단계 초기 구조로 승계한다.

    1차(빠름)가 이미 찾은 구조를 2차(표준)가 버리고 conformer 탐색부터 다시 하는
    낭비를 없앤다 — 분자당 conformer 임베딩 + DFT 재순위화(SCF 수 회)가 절약된다.
    Boltzmann 앙상블이 켜진 단계(정밀)는 여러 conformer 가중이 목적이므로 승계하지
    않고, 사용자 업로드 구조보다 이전 단계 «계산 결과» 구조를 우선한다.
    """
    if not CHAIN_GEOMETRY or stage_idx == 0:
        return c.get("geometry")
    acc = camp["stages"][stage_idx]["accuracy"]
    if presets.ACCURACY.get(acc, {}).get("ensemble"):
        return c.get("geometry")
    prev = store.get_job(c["jobs"].get(str(stage_idx - 1)))
    xyz = ((prev or {}).get("result") or {}).get("structure_xyz")
    atoms = _xyz_atoms(xyz) if xyz else None
    if atoms:
        return {"atoms": atoms, "source": f"{stage_idx}단계 결과 승계",
                "rescan": False}
    return c.get("geometry")


def _maybe_slim(job: dict):
    """캠페인 작업의 무거운 시각화 데이터 제거 — 판정·비교·3D 구조는 유지.

    전자밀도 구름은 결과당 수천 점이라 300개 캠페인이면 jobs.json 이 수십 MB로
    불어난다. 단건으로 다시 계산하면 언제든 다시 생성된다.
    """
    if not SLIM_BATCH:
        return
    res = job.get("result") or {}
    if not res or res.get("slimmed") or res.get("density_cloud") is None:
        return
    res["density_cloud"] = None
    res["slimmed"] = True
    res.setdefault("notes", []).append(
        "배치 스크리닝 용량 절약 — 전자밀도 구름을 저장하지 않았습니다 "
        "(단건으로 다시 계산하면 표시됩니다)")
    store.update_job(job["id"], {"result": res})


def _cut_score(c: dict, camp: dict, stage_idx: int):
    """깔때기 컷 순위 점수 — 안정성 여유(높을수록 생존).

    수직 전위 기반 결과(빠름 프리셋)는 EA를 과소평가해 환원 위험을 낮잡으므로,
    컷 비교에서만 환원 전위를 VERTICAL_RED_BUFFER 만큼 올려 보수적으로 본다.
    최종 판정(judge)에는 보정을 넣지 않는다 — 등급은 계산값 그대로 두고,
    낙관 가능성은 basis·note 로 표시한다.
    """
    jid = c["jobs"].get(str(stage_idx))
    job = store.get_job(jid) if jid else None
    if not job or job["status"] != "PUBLISHED" or not job.get("result"):
        return None
    desc = job["result"].get("descriptors") or {}
    red, ox = _potentials(desc)
    if red is None or ox is None:
        return None
    if (desc.get("reduction_potential_gibbs_v") is None
            and desc.get("ea_adiabatic_ev") is None):
        red = red + VERTICAL_RED_BUFFER
    # 범위 판정 대상(P0-5)은 컷 순위에서도 보수적인 끝값을 쓴다 — judge 와 같은 원칙
    rng = confsens.judgement_range(desc.get("conformer_sensitivity"))
    if rng:
        red, ox = red + rng["red"][1], ox + rng["ox"][0]
    wins = _windows(camp)
    return min(min(wins[e]["low"] - red, ox - wins[e]["high"])
               for e in camp["electrodes"])


def _advance(camp: dict):
    """캠페인 하나를 한 걸음 진행시킨다 — 제출·재시도·단계 전환·판정."""
    stage_idx = camp["stageIndex"]
    stage_key = str(stage_idx)
    alive = [c for c in camp["candidates"] if c["alive"]]
    if not alive:
        camp["status"] = "DONE"
        camp["finishedAt"] = time.time()
        _log(camp, "살아남은 후보가 없어 종료")
        return

    pending, n_done, n_failed_stage = [], 0, 0
    for c in alive:
        jid = c["jobs"].get(stage_key)
        job = store.get_job(jid) if jid else None
        if job is None:
            pending.append(c)
            continue
        if job["status"] == "PUBLISHED":
            _maybe_slim(job)
            n_done += 1
        elif job["status"] == "FAILED":
            if RESTART_ERROR in (job.get("error") or ""):
                # 서버 재시작 유탄 — 재시도 횟수를 쓰지 않고 다시 제출
                del c["jobs"][stage_key]
                pending.append(c)
                _log(camp, f"{c['name']}: 서버 재시작으로 중단 — 재제출")
            elif c["retries"].get(stage_key, 0) < 1:
                c["retries"][stage_key] = c["retries"].get(stage_key, 0) + 1
                del c["jobs"][stage_key]
                pending.append(c)
                _log(camp, f"{c['name']}: 실패({job.get('error')}) — 자동 재시도 1회")
            else:
                c["alive"] = False
                c["failed"] = True
                c["error"] = job.get("error") or "계산 실패"
                n_failed_stage += 1
                _log(camp, f"{c['name']}: 재시도 후에도 실패 — 판정 불가로 분류")
        # QUEUED / RUNNING → 기다린다

    # 실패 누적률 보호 — 조건이 잘못된 캠페인이 밤새 실패만 쌓는 것을 막는다
    n_failed_total = sum(1 for c in camp["candidates"] if c["failed"])
    if (camp["status"] == "RUNNING" and len(camp["candidates"]) >= 10
            and n_failed_total / len(camp["candidates"]) > FAIL_PAUSE_RATIO
            and not camp.get("autoPaused")):
        camp["status"] = "PAUSED"
        camp["autoPaused"] = True
        _log(camp, f"실패율 {n_failed_total}/{len(camp['candidates'])} — "
                   "임계 초과로 자동 일시정지. 조건을 확인한 뒤 재개하세요.")
        return

    # 제출 — 배치 전용 동시 실행 한도 안에서, 캐시가 있으면 계산 없이 연결
    slots = BATCH_PARALLEL - _batch_active_count()
    for c in pending:
        # 업로드 3D 구조가 있는 후보는 캐시를 쓰지 않는다 — 좌표가 다르면 다른 계산
        cached = None if c.get("geometry") else \
            _find_cached(c["canonical"], _stage_settings(camp, stage_idx))
        if cached is not None:
            c["jobs"][stage_key] = cached["id"]
            _log(camp, f"{c['name']}: 동일 조건 기존 결과 재사용 ({cached['id']})")
            continue
        if slots <= 0:
            continue
        settings = _stage_settings(camp, stage_idx)
        material = {"id": None, "name": c["name"], "abbr": "배치",
                    "smiles": c["calcSmiles"]}
        geom = _chained_geometry(c, camp, stage_idx)
        if geom:
            material["geometry"] = geom
        job = store.create_job(material, settings)
        store.update_job(job["id"], {"campaign": {
            "id": camp["id"], "name": camp["name"], "stage": stage_idx,
            "accuracy": camp["stages"][stage_idx]["accuracy"]}})
        worker.submit(job["id"], priority=worker.PRIORITY_BATCH)
        c["jobs"][stage_key] = job["id"]
        slots -= 1

    # 단계 완료 판정 — 살아 있는 모든 후보가 이 단계에서 PUBLISHED 인가
    alive = [c for c in camp["candidates"] if c["alive"]]
    done = [c for c in alive
            if (store.get_job(c["jobs"].get(stage_key)) or {}).get("status") == "PUBLISHED"]
    if not alive or len(done) < len(alive):
        return

    last_stage = stage_idx == len(camp["stages"]) - 1
    if last_stage:
        _finalize(camp)
        return

    # 깔때기 — v2.0 14.7: 고정 Top-N 대신 «임계값 + 최대 Top-N» 조합.
    # 고정 개수만 쓰면 후보가 다 좋아도 N개만 남고, 다 나빠도 N개는 통과한다.
    # 임계값이 먼저 거르고, Top-N 은 계산량 상한으로만 작동한다.
    stage_cfg = camp["stages"][stage_idx]
    keep = stage_cfg.get("keep")
    threshold = stage_cfg.get("threshold")
    scored = sorted(done, key=lambda c: (_cut_score(c, camp, stage_idx) is None,
                                         -(_cut_score(c, camp, stage_idx) or -1e9)))
    passed = scored
    if threshold is not None:
        passed = [c for c in scored
                  if (_cut_score(c, camp, stage_idx) or -1e9) >= threshold]
        _log(camp, f"{stage_idx + 1}단계 임계값 {threshold:+.2f} V 적용 — "
                   f"{len(done)}개 중 {len(passed)}개 통과")
    survivors = passed if not keep else passed[:keep]
    if keep and len(passed) > keep:
        _log(camp, f"계산량 상한(Top-{keep})에 걸려 임계값 통과 {len(passed)}개 중 "
                   f"{keep}개만 다음 단계로 보냅니다 — 나머지는 탈락이 아니라 «보류»입니다")
        for c in passed[keep:]:
            c["deferred"] = True
    survivor_ids = {c["idx"] for c in survivors}
    for c in done:
        if c["idx"] not in survivor_ids:
            c["alive"] = False
            c["cutStage"] = stage_idx
    camp["stageIndex"] = stage_idx + 1
    _log(camp, f"{stage_idx + 1}단계({camp['stages'][stage_idx]['accuracy']}) 완료 — "
               f"{len(done)}개 중 {len(survivors)}개가 "
               f"{stage_idx + 2}단계({camp['stages'][stage_idx + 1]['accuracy']})로 진출")
    if camp["stages"][stage_idx]["accuracy"] == "빠름":
        _log(camp, f"컷 순위에는 수직 전위의 낙관을 상쇄하는 환원 보수 보정 "
                   f"+{VERTICAL_RED_BUFFER} V 를 적용했습니다")
    next_acc = camp["stages"][stage_idx + 1]["accuracy"]
    if CHAIN_GEOMETRY and not presets.ACCURACY.get(next_acc, {}).get("ensemble"):
        _log(camp, f"{stage_idx + 2}단계는 이번 단계의 최적 구조에서 바로 시작합니다 "
                   "(conformer 탐색·재순위화 생략 — 계산 시간 절약)")


def _finalize(camp: dict):
    """마지막 단계 완료 — 최종 판정을 저장하고 캠페인을 닫는다."""
    stage_key = str(camp["stageIndex"])
    for c in camp["candidates"]:
        if not c["alive"]:
            continue
        job = store.get_job(c["jobs"].get(stage_key))
        if job and job["status"] == "PUBLISHED" and job.get("result"):
            desc = job["result"].get("descriptors") or {}
            c["verdict"] = judge(desc, camp["electrodes"], camp["margin_v"],
                                 windows=_windows(camp), e_abs=_e_abs(camp))
            sc = camp.get("scoring") or {}
            if sc.get("enabled"):
                c["score"] = scoring.evaluate(
                    desc, c["verdict"], camp["electrodes"],
                    weights=sc.get("weights"), preset=sc.get("preset") or "균등")
    camp["status"] = "DONE"
    camp["finishedAt"] = time.time()
    graded = [c for c in camp["candidates"] if c.get("verdict")]
    n_fit = sum(1 for c in graded if c["verdict"]["grade"] == "적합")
    n_cond = sum(1 for c in graded if c["verdict"]["grade"] == "조건부")
    _log(camp, f"캠페인 완료 — 적합 {n_fit} · 조건부 {n_cond} · "
               f"부적합 {len(graded) - n_fit - n_cond} · "
               f"판정 불가 {sum(1 for c in camp['candidates'] if c['failed'])}")


def _promote_queued() -> bool:
    """실행 중 캠페인이 없으면 대기열의 가장 오래된 캠페인을 시작한다 (호출자가 _lock 보유)."""
    if any(c["status"] == "RUNNING" for c in _campaigns.values()):
        return False
    queued = sorted((c for c in _campaigns.values() if c["status"] == "QUEUED"),
                    key=lambda c: c["createdAt"])
    if not queued:
        return False
    camp = queued[0]
    camp["status"] = "RUNNING"
    _log(camp, "앞 캠페인이 끝나 대기열에서 자동 시작")
    return True


_TICK_S = 3.0
_orc_started = False
_orc_lock = threading.Lock()


def _orchestrator_loop():
    while True:
        try:
            with _lock:
                _load()
                running = [c for c in _campaigns.values() if c["status"] == "RUNNING"]
                for camp in running:
                    before = json.dumps(camp, sort_keys=True, ensure_ascii=False,
                                        default=str)
                    _advance(camp)
                    if json.dumps(camp, sort_keys=True, ensure_ascii=False,
                                  default=str) != before:
                        _persist()
                if _promote_queued():
                    _persist()
        except Exception:  # noqa: BLE001 — 오케스트레이터는 죽지 않아야 함
            pass
        time.sleep(_TICK_S)


def ensure_started():
    global _orc_started
    with _orc_lock:
        if not _orc_started:
            threading.Thread(target=_orchestrator_loop, daemon=True,
                             name="screening-orchestrator").start()
            _orc_started = True


# ---------------------------------------------------------------- 조회용 뷰
def _candidate_view(c: dict, camp: dict) -> dict:
    stage_idx = camp["stageIndex"]
    jid = c["jobs"].get(str(stage_idx))
    job = store.get_job(jid) if jid else None
    if c["failed"]:
        state, detail = "판정 불가", c.get("error")
    elif c["cutStage"] is not None:
        state = "탈락"
        detail = f"{c['cutStage'] + 1}단계에서 순위 미달"
    elif camp["status"] in ("DONE", "CANCELLED"):
        state, detail = ("완료", None) if c.get("verdict") else ("중단", None)
    elif job is None:
        state, detail = "대기", None
    else:
        state = {"QUEUED": "대기", "RUNNING": "계산 중",
                 "PUBLISHED": "완료", "FAILED": "재시도 대기"}.get(job["status"], "대기")
        detail = job.get("stage") if job["status"] == "RUNNING" else None

    # 판정 — 저장본을 쓰지 않고 항상 다시 판정한다 (binder_rank 와 같은 원칙):
    # 순수 함수라 비용이 없고, 판정 로직이 개선되면(예: LUMO 불일치 강등)
    # 이미 끝난 캠페인의 화면·집계·CSV 에도 즉시 반영된다. 마지막 완료 단계 기준.
    verdict, provisional, desc = None, False, None
    for si in range(len(camp["stages"]) - 1, -1, -1):
        j = store.get_job(c["jobs"].get(str(si)))
        if j and j["status"] == "PUBLISHED" and j.get("result"):
            desc = j["result"].get("descriptors") or {}
            verdict = judge(desc, camp["electrodes"], camp["margin_v"],
                            windows=_windows(camp), e_abs=_e_abs(camp))
            provisional = (si < len(camp["stages"]) - 1
                           or camp["status"] not in ("DONE", "CANCELLED"))
            break
    if verdict is None:
        verdict = c.get("verdict")   # 작업이 지워진 경우에만 저장본 폴백
    score = c.get("score")
    sc_cfg = camp.get("scoring") or {}
    if sc_cfg.get("enabled") and desc is not None and verdict is not None:
        score = scoring.evaluate(desc, verdict, camp["electrodes"],
                                 weights=sc_cfg.get("weights"),
                                 preset=sc_cfg.get("preset") or "균등")
    # 계산 검증 등급(PASS/REVIEW/FAIL) — 마지막 작업의 monitor 검증 보고서.
    # 판정(적합/부적합)과는 별개다: 프로세스가 끝났다고 계산이 정확한 것은 아니다.
    last_job = None
    for si in range(len(camp["stages"]) - 1, -1, -1):
        last_job = store.get_job(c["jobs"].get(str(si)))
        if last_job:
            break
    val = (last_job or {}).get("validation") if last_job else None
    mon = (last_job or {}).get("monitor") or {}
    return {"idx": c["idx"], "name": c["name"], "smiles": c["smiles"],
            "calcSmiles": c["calcSmiles"], "atoms": c["atoms"],
            "state": state, "detail": detail,
            "alive": c["alive"], "failed": c["failed"], "cutStage": c["cutStage"],
            "jobs": c["jobs"], "verdict": verdict, "provisional": provisional,
            "score": score, "fgroups": c.get("fgroups") or [],
            "progress": (job or {}).get("progress") if job else None,
            "validation": ({"grade": val.get("grade"), "summary": val.get("summary")}
                           if val else None),
            "monitorJob": last_job["id"] if last_job else None,
            "scf": (mon.get("scf") or {}).get("cycle") if job and job["status"] == "RUNNING" else None,
            "anomalies": sum(1 for a in (mon.get("anomalies") or [])
                             if a.get("severity") in ("WARN", "ERROR", "FATAL"))}


def _overall_pct(camp: dict) -> int:
    """캠페인 전체 진행률(%) — 단계별 예상 비용(초/분자)으로 가중한다.

    지나온 단계는 완료로, 현재 단계는 완료/생존 비율로, 앞으로의 단계는
    통과 수 기준 예정량으로 집계한다.
    """
    if camp["status"] == "DONE":
        return 100
    if camp["status"] == "CANCELLED":
        return 0
    stage_idx = camp["stageIndex"]
    n_alive = sum(1 for c in camp["candidates"] if c["alive"])
    total = done = 0.0
    planned = n_alive
    for si, st in enumerate(camp["stages"]):
        w = ESTIMATE_S.get(st["accuracy"], 1200)
        if si < stage_idx:
            n = sum(1 for c in camp["candidates"] if str(si) in c["jobs"])
            total += n * w
            done += n * w
        elif si == stage_idx:
            n_done = sum(1 for c in camp["candidates"]
                         if (store.get_job(c["jobs"].get(str(si))) or {})
                         .get("status") == "PUBLISHED")
            total += n_alive * w
            done += n_done * w
            planned = min(n_alive, st.get("keep") or n_alive)
        else:
            total += planned * w
            planned = min(planned, st.get("keep") or planned)
    return int(round(100 * done / total)) if total else 0


def campaign_view(camp: dict) -> dict:
    """진행률·판정을 계산해 붙인 캠페인 상세 — API 응답용."""
    stage_idx = camp["stageIndex"]
    stages = []
    for si, st in enumerate(camp["stages"]):
        entered = [c for c in camp["candidates"]
                   if str(si) in c["jobs"] or (c["alive"] and si == stage_idx)
                   or (c["cutStage"] is not None and si <= c["cutStage"])]
        jobs = [store.get_job(c["jobs"].get(str(si))) for c in camp["candidates"]]
        jobs = [j for j in jobs if j]
        done = sum(1 for c in camp["candidates"]
                   if (store.get_job(c["jobs"].get(str(si))) or {}).get("status") == "PUBLISHED")
        stages.append({"accuracy": st["accuracy"], "keep": st.get("keep"),
                       "n_entered": len(entered) if si <= stage_idx else 0,
                       "n_done": done,
                       "state": ("완료" if si < stage_idx
                                 or (si == stage_idx and camp["status"] == "DONE")
                                 else "대기" if si > stage_idx
                                 else {"RUNNING": "진행 중", "PAUSED": "일시정지",
                                       "QUEUED": "대기", "CANCELLED": "중단"}
                                 .get(camp["status"], "진행 중"))})

    cands = [_candidate_view(c, camp) for c in camp["candidates"]]
    # 순위 — 판정 가능한 후보를 안정성 여유 내림차순으로
    if (camp.get("scoring") or {}).get("enabled"):
        # 5대 Score 캠페인 — Hard Filter 탈락은 뒤로, 나머지는 총점 내림차순
        ranked = sorted([c for c in cands
                         if c.get("score") and c["score"]["total"] is not None],
                        key=lambda c: (c["score"]["hard_fail"], -c["score"]["total"]))
    else:
        ranked = sorted([c for c in cands if c["verdict"] and
                         c["verdict"]["worst_margin_v"] is not None],
                        key=lambda c: -c["verdict"]["worst_margin_v"])
    for rank, c in enumerate(ranked, start=1):
        c["rank"] = rank

    n_alive = sum(1 for c in camp["candidates"] if c["alive"])
    cur_done = stages[stage_idx]["n_done"] if stage_idx < len(stages) else 0
    # 남은 시간 추정 — 이 단계에서 실측된 계산 시간이 있으면 그것을 쓴다
    walls = []
    for c in camp["candidates"]:
        j = store.get_job(c["jobs"].get(str(stage_idx)))
        if j and j["status"] == "PUBLISHED" and j.get("result"):
            w = j["result"].get("wall_time_s")
            if isinstance(w, (int, float)):
                walls.append(w)
    per_mol = (sum(walls) / len(walls)) if walls else \
        ESTIMATE_S.get(camp["stages"][stage_idx]["accuracy"] if stage_idx < len(camp["stages"]) else "표준", 1200)
    remaining = max(0, n_alive - cur_done)
    eta_s = None
    if camp["status"] == "RUNNING" and remaining:
        eta_s = int(remaining * per_mol / BATCH_PARALLEL)
        # 뒤 단계 추정도 더한다
        for si in range(stage_idx + 1, len(camp["stages"])):
            n = camp["stages"][si - 1].get("keep") or remaining
            eta_s += int(n * ESTIMATE_S.get(camp["stages"][si]["accuracy"], 1200)
                         / BATCH_PARALLEL)

    counts = {
        "total": len(camp["candidates"]),
        "alive": n_alive,
        "fit": sum(1 for c in cands if c["verdict"] and not c["provisional"]
                   and c["verdict"]["grade"] == "적합"),
        "conditional": sum(1 for c in cands if c["verdict"] and not c["provisional"]
                           and c["verdict"]["grade"] == "조건부"),
        "unfit": sum(1 for c in cands if c["verdict"] and not c["provisional"]
                     and c["verdict"]["grade"] == "부적합"),
        "failed": sum(1 for c in cands if c["failed"]),
        "cut": sum(1 for c in cands if c["cutStage"] is not None),
    }
    return {**{k: v for k, v in camp.items() if k != "candidates"},
            "stages": stages, "candidates": cands, "counts": counts,
            "eta_s": eta_s, "batch_parallel": BATCH_PARALLEL,
            "overall_pct": _overall_pct(camp),
            "scoring": camp.get("scoring"),
            "protocol": camp.get("protocol"),
            "thermo_note": THERMO_NOTE}


def reverify_lumo(camp: dict) -> dict:
    """LUMO 불일치 후보만 모아 표준 정확도로 다시 확인하는 후속 캠페인.

    같은 전극·마진·용매·설정을 그대로 쓰고 정확도만 «표준»(단열·ΔG 전위)으로
    올린다 — 수직 EA 와 LUMO 추정이 상반된 후보의 확정 판정이 목적이다.
    다른 캠페인이 돌고 있으면 대기열로 들어간다.
    """
    view = campaign_view(camp)
    targets = [c for c in view["candidates"]
               if (c.get("verdict") or {}).get("lumo_check")]
    if not targets:
        raise ValueError("LUMO 불일치 후보가 없습니다.")
    rows = []
    for t in targets:
        r = parse_candidates(f"{t['name']},{t['smiles']}")["rows"]
        if r and r[0]["ok"]:
            rows.append(r[0])
    if not rows:
        raise ValueError("재검증할 후보를 해석하지 못했습니다.")
    settings = dict(camp["settings"])
    settings["accuracy"] = "표준"
    return create_campaign(
        name=f"{camp['name']} — LUMO 재검증(표준)", candidates=rows,
        electrodes=camp["electrodes"], margin_v=camp["margin_v"],
        stages=[{"accuracy": "표준"}], settings=settings,
        custom_electrodes=camp.get("customElectrodes"),
        scoring_cfg=camp.get("scoring"))


def campaign_summary(camp: dict) -> dict:
    stage_idx = camp["stageIndex"]
    stage_key = str(stage_idx)
    n_done_stage = sum(
        1 for c in camp["candidates"]
        if (store.get_job(c["jobs"].get(stage_key)) or {}).get("status") == "PUBLISHED")
    return {"id": camp["id"], "name": camp["name"], "status": camp["status"],
            "createdAt": camp["createdAt"], "finishedAt": camp["finishedAt"],
            "autoPaused": camp.get("autoPaused", False),
            "electrodes": camp["electrodes"], "margin_v": camp["margin_v"],
            "stages": [st["accuracy"] for st in camp["stages"]],
            "stageIndex": stage_idx,
            "n_candidates": len(camp["candidates"]),
            "n_alive": sum(1 for c in camp["candidates"] if c["alive"]),
            "n_failed": sum(1 for c in camp["candidates"] if c["failed"]),
            "n_done_stage": n_done_stage,
            "overall_pct": _overall_pct(camp)}


def estimate(n_candidates: int, stages: list[dict]) -> dict:
    """실행 확인 화면용 — 단계별 대상 수와 예상 소요."""
    out, n = [], n_candidates
    total_s = 0
    for st in stages:
        per = ESTIMATE_S.get(st["accuracy"], 1200)
        sec = int(n * per / BATCH_PARALLEL)
        out.append({"accuracy": st["accuracy"], "n": n, "per_mol_s": per,
                    "stage_s": sec, "keep": st.get("keep"),
                    "threshold": st.get("threshold")})
        total_s += sec
        # 임계값 통과 수는 계산 전에는 알 수 없다 — 예상치는 Top-N 상한으로 잡는다
        n = min(n, st.get("keep") or n)
    return {"stages": out, "total_s": total_s, "batch_parallel": BATCH_PARALLEL}
